"""Investment committee report generator (Markdown and Excel).

Produces comprehensive stress-testing reports for frontier market deals,
including devaluation probabilities, return distributions, hurdle risk,
required local IRR, hedging comparisons, risk drivers, and governance disclaimers.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import numpy as np
import pandas as pd

from parity.deal import DealAnalysisResult
from parity.devaluation import DevaluationModel
from parity.fxpaths import SimulationResult
from parity.hedging import HedgeAnalysisResult
from parity.schemas import DealConfig


def compute_top_risk_drivers(
    features: Dict[str, float],
    model: Optional[DevaluationModel] = None,
    n_top: int = 3,
) -> List[Tuple[str, float, str]]:
    """Identify the top macro risk drivers contributing to devaluation log-odds.

    Parameters
    ----------
    features : Dict[str, float]
        Country macro features.
    model : Optional[DevaluationModel], default None
        Devaluation model instance. If None, default model is initialized.
    n_top : int, default 3
        Number of top drivers to return.

    Returns
    -------
    List[Tuple[str, float, str]]
        List of (feature_name, log_odds_contribution, description) tuples.
    """
    if model is None:
        model = DevaluationModel(horizon=12)

    contributions: List[Tuple[str, float, str]] = []
    descriptions = {
        "parallel_premium": "Parallel market exchange rate spread over official peg",
        "reserves_cover": "Depletion of foreign exchange reserves cover",
        "reserves_months_imports": "Depletion of foreign exchange reserves cover",
        "reserves_change_3m": "Rapid 3-month drain in central bank liquid reserves",
        "reserves_change_3m_pct": "Rapid 3-month drain in central bank liquid reserves",
        "ca_pct_gdp": "Structural current account deficit",
        "current_account_pct_gdp": "Structural current account deficit",
        "rer_misalignment": "Real effective exchange rate (REER) overvaluation",
        "rate_differential": "Domestic policy rate carry defense against USD",
        "imf_program": "Absence of IMF program external financing backstop",
        "imf_review_off_track": "Derailed or stalled IMF program performance review",
    }

    alias_map = {
        "reserves_months_imports": "reserves_cover",
        "reserves_change_3m_pct": "reserves_change_3m",
        "current_account_pct_gdp": "ca_pct_gdp",
    }

    for feat_name, val in features.items():
        model_feat = alias_map.get(feat_name, feat_name)
        if model_feat in model.coef_:
            coef = model.coef_[model_feat]
            # Contribution to log odds: feature_value * coefficient
            # Higher positive contribution increases devaluation probability
            contrib = float(val) * float(coef)
            desc = descriptions.get(feat_name, descriptions.get(model_feat, f"Macro indicator: {feat_name}"))
            contributions.append((feat_name, contrib, desc))

    # Sort descending by contribution to devaluation risk
    contributions.sort(key=lambda x: x[1], reverse=True)
    return contributions[:n_top]


def generate_markdown_report(
    deal_res: DealAnalysisResult,
    deal_config: DealConfig,
    dev_prob_12m: float,
    dev_prob_24m: float,
    risk_drivers: List[Tuple[str, float, str]],
    hedge_res: Optional[HedgeAnalysisResult] = None,
    comparison_df: Optional[pd.DataFrame] = None,
    as_of_date: Optional[Union[str, dt.date]] = None,
    country: Optional[str] = None,
) -> str:
    """Generate a clean, structured Markdown report for an Investment Committee.

    Parameters
    ----------
    deal_res : DealAnalysisResult
        Simulated deal returns and risk metrics.
    deal_config : DealConfig
        Deal configuration parameters.
    dev_prob_12m : float
        12-month devaluation probability (0.0 to 1.0).
    dev_prob_24m : float
        24-month devaluation probability (0.0 to 1.0).
    risk_drivers : List[Tuple[str, float, str]]
        Top risk drivers from compute_top_risk_drivers.
    hedge_res : Optional[HedgeAnalysisResult], default None
        Hedging comparison results if forwards or CIP proxy available.
    comparison_df : Optional[pd.DataFrame], default None
        Scenario switchboard comparison table.
    as_of_date : Optional[Union[str, dt.date]], default None
        As-of date for macro data.
    country : Optional[str], default None
        Target country name.

    Returns
    -------
    str
        Markdown formatted report string.
    """
    country_name = country or deal_config.country
    date_str = str(as_of_date or deal_config.close_date)

    md: List[str] = []

    # Title & Metadata
    md.append(f"# Parity Stress-Test Report: {country_name}")
    md.append(f"**As-of Date:** {date_str} | **Deal Close Date:** {deal_config.close_date}")
    md.append(f"**Target Hurdle (USD IRR):** {deal_config.hurdle_usd_irr * 100:.1f}%\n")

    # Data Freshness Warning
    md.append("> [!WARNING]")
    md.append(f"> **Data Freshness Warning:** Macro data evaluated as of **{date_str}**.")
    md.append("> Frontier market conditions (parallel spreads, central bank reserves) can change rapidly.")
    md.append("> Always verify parallel market quotes and official FX window liquidity prior to capital deployment.\n")

    # Section 1: Executive Summary & Return Metrics
    md.append("## 1. Executive Return Summary")
    md.append(
        "Headline local-currency returns versus simulated hard-currency (USD) returns "
        "accounting for simulated FX trajectories, parallel-market repatriation, and withholding taxes:\n"
    )

    md.append("| Metric | LCU Headline | USD Expected (Mean) | USD Median | USD P10 (Stress) | USD P90 (Upside) |")
    md.append("| :--- | :---: | :---: | :---: | :---: | :---: |")
    md.append(
        f"| **Annualized IRR** | **{deal_res.headline_lcu_irr * 100:.1f}%** | "
        f"**{deal_res.usd_irr_mean * 100:.1f}%** | "
        f"{deal_res.usd_irr_median * 100:.1f}% | "
        f"{deal_res.usd_irr_p10 * 100:.1f}% | "
        f"{deal_res.usd_irr_p90 * 100:.1f}% |"
    )
    md.append(
        f"| **MOIC** | **{deal_res.headline_lcu_moic:.2f}x** | "
        f"**{deal_res.usd_moic_mean:.2f}x** | "
        f"{deal_res.usd_moic_median:.2f}x | "
        f"{deal_res.usd_moic_p10:.2f}x | "
        f"{deal_res.usd_moic_p90:.2f}x |\n"
    )

    # Key Takeaways Bulleted
    fx_drag_bps = deal_res.expected_fx_drag * 10000
    p_shortfall = deal_res.p_below_hurdle * 100
    md.append(f"- **Expected FX Drag:** `{deal_res.expected_fx_drag * 100:.2f}%` ({fx_drag_bps:.0f} bps) annual reduction from local headline return to expected USD return.")
    md.append(f"- **Hurdle Shortfall Risk:** `{p_shortfall:.1f}%` probability that USD IRR falls below the {deal_config.hurdle_usd_irr * 100:.1f}% hurdle.")
    md.append(f"- **Required Local-Currency IRR (Breakeven):**")
    md.append(f"  - To achieve {deal_config.hurdle_usd_irr * 100:.1f}% USD IRR at **50% confidence**: **`{deal_res.breakeven_lcu_irr_50 * 100:.1f}%`** LCU IRR.")
    md.append(f"  - To achieve {deal_config.hurdle_usd_irr * 100:.1f}% USD IRR at **75% confidence**: **`{deal_res.breakeven_lcu_irr_75 * 100:.1f}%`** LCU IRR.\n")

    # Section 2: Devaluation Hazard & Risk Drivers
    md.append("## 2. Devaluation Hazard Assessment")
    md.append(f"- **12-Month Devaluation Probability:** **`{dev_prob_12m * 100:.1f}%`**")
    md.append(f"- **24-Month Devaluation Probability:** **`{dev_prob_24m * 100:.1f}%`**\n")

    md.append("### Top Drivers of Devaluation Risk")
    md.append("| Rank | Macro Risk Indicator | Log-Odds Contribution | Mechanism / Channel |")
    md.append("| :---: | :--- | :---: | :--- |")
    for i, (feat, contrib, desc) in enumerate(risk_drivers, 1):
        md.append(f"| {i} | `{feat}` | `+{contrib:.2f}` | {desc} |")
    md.append("")

    # Section 3: Hedging Analysis
    if hedge_res is not None:
        md.append("## 3. Hedging Comparison & Cost of Carry")
        proxy_note = " *(Covered Interest Parity Proxy)*" if hedge_res.is_cip_proxy else " *(Market Forwards)*"
        md.append(f"Hedging basis: **{hedge_res.forward_source}{proxy_note}**\n")

        md.append("| Metric | Unhedged Deal | Fully Hedged Deal | Variance / Drag |")
        md.append("| :--- | :---: | :---: | :---: |")
        hedged_irr_str = f"{hedge_res.hedged_usd_irr * 100:.1f}%" if not np.isnan(hedge_res.hedged_usd_irr) else "N/A"
        md.append(f"| **USD IRR** | {hedge_res.unhedged_usd_irr_mean * 100:.1f}% (Expected) | {hedged_irr_str} | -{hedge_res.hedge_cost_drag * 100:.1f}% |")
        md.append(f"| **USD MOIC** | {deal_res.usd_moic_mean:.2f}x | {hedge_res.hedged_usd_moic:.2f}x | -{deal_res.usd_moic_mean - hedge_res.hedged_usd_moic:.2f}x |")
        md.append(f"| **Breakeven Annual Depreciation** | — | **{hedge_res.breakeven_annual_depreciation * 100:.1f}% / yr** | Level where hedging breaks even |")
        md.append(f"| **P(Hedge Outperforms Unhedged)** | — | **{hedge_res.p_hedge_outperforms * 100:.1f}%** | Odds that unhedged suffers worse drag than forward |")
        md.append("")

    # Section 4: Scenario Switchboard Table (if provided)
    if comparison_df is not None and not comparison_df.empty:
        md.append("## 4. Scenario Stress-Testing Switchboard")
        md.append("Comparative deal performance across macro stress scenarios:\n")
        md.append("| Scenario | Description | Mean USD IRR | Median USD IRR | P10 (Stress) | P(USD IRR < Hurdle) |")
        md.append("| :--- | :--- | :---: | :---: | :---: | :---: |")
        for _, row in comparison_df.iterrows():
            scen_name = row.get("scenario", "")
            desc = row.get("description", "")
            mean_irr = float(row.get("usd_irr_mean", np.nan)) * 100
            med_irr = float(row.get("usd_irr_median", np.nan)) * 100
            p10_irr = float(row.get("usd_irr_p10", np.nan)) * 100
            p_short = float(row.get("p_below_hurdle", np.nan)) * 100
            md.append(f"| **{scen_name}** | {desc} | {mean_irr:.1f}% | {med_irr:.1f}% | {p10_irr:.1f}% | {p_short:.1f}% |")
        md.append("")

    # Section 5: Repatriation & Frictions
    md.append("## 5. Repatriation & Capital Control Frictions")
    md.append(f"- **Repatriation Mechanism:** `{deal_config.repatriation_rate}` (Blend Weight: {deal_config.repatriation_blend_weight * 100:.0f}%)")
    md.append(f"- **Blocked Funds / Trapped Cash Haircut:** `{deal_config.blocked_funds_haircut_pct * 100:.1f}%`")
    md.append(f"- **Cross-Border Withholding Tax (WHT):** `{deal_config.withholding_tax_pct * 100:.1f}%`\n")

    # Section 6: Governance & Model Limitations
    md.append("## 6. Model Limitations & Governance Disclaimers")
    md.append("> [!IMPORTANT]")
    md.append("> **Model Limitations:**")
    md.append("> 1. **Sparse Historical Devaluation Episodes:** Frontier market discrete devaluations occur infrequently; model priors govern behavior unless data history is extensive.")
    md.append("> 2. **Realized Repatriation Rates vs Official Rates:** Official exchange rates frequently do not reflect the effective rate realized upon repatriation due to backlog queues, parallel premiums, and allocation rationing.")
    md.append("> 3. **Policy Discretion & Timing:** Devaluations are political and central bank policy decisions; the model computes probability hazards, not calendar timing.")
    md.append("> 4. **Forward / NDF Market Illiquidity:** Market forward quotes in frontier currencies are often illiquid or indicative; CIP proxy rates reflect theoretical carrying cost rather than executable liquidity.")
    md.append("> 5. **No Investment Advice:** Parity outputs are quantitative stress-testing estimates provided on an 'as-is' basis under Apache 2.0 without warranty. Not intended as investment advice or credit ratings.\n")

    return "\n".join(md)


def generate_excel_report(
    deal_res: DealAnalysisResult,
    deal_config: DealConfig,
    dev_prob_12m: float,
    dev_prob_24m: float,
    risk_drivers: List[Tuple[str, float, str]],
    deal_df: pd.DataFrame,
    hedge_res: Optional[HedgeAnalysisResult] = None,
    comparison_df: Optional[pd.DataFrame] = None,
    sim_result: Optional[SimulationResult] = None,
    as_of_date: Optional[Union[str, dt.date]] = None,
    country: Optional[str] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """Generate an institutional Excel workbook with executive formatting, charts, and conditional colors.

    Tabs created:
    1. Executive Summary: KPI blocks, Return metrics, Hurdle comparison, Hedge summary, Risk drivers.
    2. Scenario Analysis: Switchboard table with conditional formatting & BarChart.
    3. FX Trajectory: Time-series percentiles of simulated FX rate & LineChart.
    4. Cash Flows: Schedule of local and converted USD cash flows.
    5. Governance & Limitations: Institutional notes and methodology.

    Parameters
    ----------
    deal_res : DealAnalysisResult
        Simulated deal returns.
    deal_config : DealConfig
        Deal configuration.
    dev_prob_12m : float
        12-month devaluation probability.
    dev_prob_24m : float
        24-month devaluation probability.
    risk_drivers : List[Tuple[str, float, str]]
        Top risk drivers.
    deal_df : pd.DataFrame
        Deal cash flow schedule.
    hedge_res : Optional[HedgeAnalysisResult], default None
        Hedge comparison result.
    comparison_df : Optional[pd.DataFrame], default None
        Scenario comparison table.
    sim_result : Optional[SimulationResult], default None
        Simulated FX paths.
    as_of_date : Optional[Union[str, dt.date]], default None
        As-of date.
    country : Optional[str], default None
        Country name.
    output_path : Optional[Union[str, Path]], default None
        Path to save .xlsx file.

    Returns
    -------
    Path
        Path to the saved Excel workbook.
    """
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    country_name = country or deal_config.country
    date_str = str(as_of_date or deal_config.close_date)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling Palettes (Institutional Deep Navy / Slate)
    NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    SUB_HEADER = PatternFill(start_color="335384", end_color="335384", fill_type="solid")
    CARD_BG = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    GREEN_FILL = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    RED_FILL = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
    AMBER_FILL = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")

    FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="1B365D")
    FONT_SECTION = Font(name="Calibri", size=12, bold=True, color="1B365D")
    FONT_HEADER_WHITE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_REGULAR = Font(name="Calibri", size=10)
    FONT_BOLD = Font(name="Calibri", size=10, bold=True)
    FONT_GREEN = Font(name="Calibri", size=11, bold=True, color="0D652D")
    FONT_RED = Font(name="Calibri", size=11, bold=True, color="C5221F")
    FONT_CARD_VAL = Font(name="Calibri", size=14, bold=True, color="1B365D")
    FONT_CARD_LBL = Font(name="Calibri", size=9, bold=False, color="475569")

    THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
    BORDER_ALL = Border(
        left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE
    )
    BORDER_TOP_BOTTOM = Border(top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

    # -------------------------------------------------------------
    # TAB 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title
    ws_sum["B2"] = f"Parity Return Stress-Testing Model — {country_name}"
    ws_sum["B2"].font = FONT_TITLE
    ws_sum["B3"] = f"Deal Close: {deal_config.close_date} | Valuation As-Of: {date_str} | Target Hurdle: {deal_config.hurdle_usd_irr * 100:.1f}%"
    ws_sum["B3"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    # KPI Summary Cards in Row 5-7
    cards = [
        ("Headline LCU IRR", f"{deal_res.headline_lcu_irr * 100:.1f}%", "Local currency basis"),
        ("Expected USD IRR", f"{deal_res.usd_irr_mean * 100:.1f}%", f"Median: {deal_res.usd_irr_median * 100:.1f}%"),
        ("Hurdle Shortfall", f"{deal_res.p_below_hurdle * 100:.1f}%", f"P(USD < {deal_config.hurdle_usd_irr*100:.0f}%)"),
        ("Expected FX Drag", f"{deal_res.expected_fx_drag * 100:.1f}%", f"{deal_res.expected_fx_drag * 10000:.0f} bps drag"),
        ("Breakeven LCU (75%)", f"{deal_res.breakeven_lcu_irr_75 * 100:.1f}%", "For target hurdle"),
    ]

    col_start = 2
    for title, val, sub in cards:
        c_letter = get_column_letter(col_start)
        ws_sum[f"{c_letter}5"] = title.upper()
        ws_sum[f"{c_letter}5"].font = FONT_CARD_LBL
        ws_sum[f"{c_letter}5"].fill = CARD_BG
        ws_sum[f"{c_letter}5"].alignment = Alignment(horizontal="center")

        ws_sum[f"{c_letter}6"] = val
        ws_sum[f"{c_letter}6"].font = FONT_CARD_VAL
        ws_sum[f"{c_letter}6"].fill = CARD_BG
        ws_sum[f"{c_letter}6"].alignment = Alignment(horizontal="center")

        ws_sum[f"{c_letter}7"] = sub
        ws_sum[f"{c_letter}7"].font = FONT_CARD_LBL
        ws_sum[f"{c_letter}7"].fill = CARD_BG
        ws_sum[f"{c_letter}7"].alignment = Alignment(horizontal="center")

        for r in range(5, 8):
            ws_sum[f"{c_letter}{r}"].border = BORDER_ALL
        col_start += 1

    # Return Distributions Table (Row 9)
    ws_sum["B9"] = "RETURN METRICS & PERCENTILES"
    ws_sum["B9"].font = FONT_SECTION

    headers_ret = ["Metric", "LCU Headline", "USD Expected (Mean)", "USD Median", "USD P10 (Stress)", "USD P90 (Upside)"]
    for j, h in enumerate(headers_ret, start=2):
        cell = ws_sum.cell(row=10, column=j, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER_WHITE
        cell.alignment = Alignment(horizontal="center")

    ret_data = [
        ("Annualized IRR", deal_res.headline_lcu_irr, deal_res.usd_irr_mean, deal_res.usd_irr_median, deal_res.usd_irr_p10, deal_res.usd_irr_p90, "0.0%"),
        ("MOIC (Multiple)", deal_res.headline_lcu_moic, deal_res.usd_moic_mean, deal_res.usd_moic_median, deal_res.usd_moic_p10, deal_res.usd_moic_p90, '0.00"x"'),
    ]

    for i, row in enumerate(ret_data, start=11):
        ws_sum.cell(row=i, column=2, value=row[0]).font = FONT_BOLD
        fmt = row[6]
        for col_idx in range(3, 8):
            c = ws_sum.cell(row=i, column=col_idx, value=row[col_idx - 2])
            c.number_format = fmt
            c.font = FONT_REGULAR
            c.alignment = Alignment(horizontal="right")
            c.border = BORDER_ALL

            # Green / Red formatting for returns exceeding hurdle
            if "IRR" in row[0]:
                val_num = float(row[col_idx - 2])
                if val_num >= deal_config.hurdle_usd_irr:
                    c.fill = GREEN_FILL
                    c.font = FONT_GREEN
                elif val_num < 0:
                    c.fill = RED_FILL
                    c.font = FONT_RED
                else:
                    c.fill = AMBER_FILL

    # Devaluation Hazard & Risk Drivers (Row 14)
    ws_sum["B14"] = "DEVALUATION HAZARD & TOP MACRO DRIVERS"
    ws_sum["B14"].font = FONT_SECTION

    ws_sum["B15"] = "12-Month Devaluation Probability"
    ws_sum["C15"] = dev_prob_12m
    ws_sum["C15"].number_format = "0.0%"
    ws_sum["C15"].font = FONT_RED if dev_prob_12m > 0.3 else FONT_BOLD

    ws_sum["B16"] = "24-Month Devaluation Probability"
    ws_sum["C16"] = dev_prob_24m
    ws_sum["C16"].number_format = "0.0%"
    ws_sum["C16"].font = FONT_RED if dev_prob_24m > 0.4 else FONT_BOLD

    # Drivers table
    driver_headers = ["Rank", "Risk Indicator", "Log-Odds Contribution", "Macro Mechanism"]
    for j, h in enumerate(driver_headers, start=2):
        c = ws_sum.cell(row=18, column=j, value=h)
        c.fill = SUB_HEADER
        c.font = FONT_HEADER_WHITE
        c.alignment = Alignment(horizontal="center")

    for i, (feat, contrib, desc) in enumerate(risk_drivers, start=19):
        ws_sum.cell(row=i, column=2, value=i - 18).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=i, column=3, value=feat).font = FONT_BOLD
        c_contrib = ws_sum.cell(row=i, column=4, value=contrib)
        c_contrib.number_format = "+0.00;-0.00;0.00"
        c_contrib.alignment = Alignment(horizontal="right")
        ws_sum.cell(row=i, column=5, value=desc).font = FONT_REGULAR
        for col_idx in range(2, 6):
            ws_sum.cell(row=i, column=col_idx).border = BORDER_ALL

    # Hedging Table (Row 23)
    ws_sum["B23"] = "HEDGING ANALYSIS & BREAKEWENS"
    ws_sum["B23"].font = FONT_SECTION

    if hedge_res is not None:
        hedge_headers = ["Hedging Metric", "Value", "Notes"]
        for j, h in enumerate(hedge_headers, start=2):
            c = ws_sum.cell(row=24, column=j, value=h)
            c.fill = NAVY_HEADER
            c.font = FONT_HEADER_WHITE

        h_rows = [
            ("Forward Pricing Source", hedge_res.forward_source, "Quoted forwards or CIP proxy"),
            ("Unhedged Expected USD IRR", hedge_res.unhedged_usd_irr_mean, "0.0%"),
            ("Fully Hedged USD IRR", hedge_res.hedged_usd_irr, "0.0%"),
            ("Annual Hedge Drag", hedge_res.hedge_cost_drag, "0.0%"),
            ("Breakeven Annual Depreciation", hedge_res.breakeven_annual_depreciation, "0.0%"),
            ("Probability Hedge Outperforms", hedge_res.p_hedge_outperforms, "0.0%"),
        ]
        for i, (m, v, n) in enumerate(h_rows, start=25):
            ws_sum.cell(row=i, column=2, value=m).font = FONT_BOLD
            val_cell = ws_sum.cell(row=i, column=3, value=v)
            if isinstance(v, float):
                val_cell.number_format = n
                val_cell.alignment = Alignment(horizontal="right")
            else:
                val_cell.alignment = Alignment(horizontal="left")
            ws_sum.cell(row=i, column=4, value=n).font = FONT_REGULAR
            for col_idx in range(2, 5):
                ws_sum.cell(row=i, column=col_idx).border = BORDER_ALL

    # -------------------------------------------------------------
    # TAB 2: SCENARIO ANALYSIS
    # -------------------------------------------------------------
    ws_scen = wb.create_sheet(title="Scenario Stress-Testing")
    ws_scen.views.sheetView[0].showGridLines = True

    ws_scen["B2"] = "Scenario Stress-Testing Switchboard"
    ws_scen["B2"].font = FONT_TITLE
    ws_scen["B3"] = f"Comparative deal performance across defined macro stress paths | Target Hurdle = {deal_config.hurdle_usd_irr * 100:.1f}%"
    ws_scen["B3"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    scen_cols = [
        "Scenario",
        "Description",
        "Headline LCU IRR",
        "Mean USD IRR",
        "Median USD IRR",
        "P10 USD IRR",
        "P90 USD IRR",
        "USD MOIC",
        "P(Below Hurdle)",
        "Expected FX Drag",
    ]
    for j, h in enumerate(scen_cols, start=2):
        c = ws_scen.cell(row=5, column=j, value=h)
        c.fill = NAVY_HEADER
        c.font = FONT_HEADER_WHITE
        c.alignment = Alignment(horizontal="center")

    scen_data_df = comparison_df if (comparison_df is not None and not comparison_df.empty) else pd.DataFrame([
        {
            "scenario": "base",
            "description": "Baseline Monte Carlo",
            "headline_lcu_irr": deal_res.headline_lcu_irr,
            "usd_irr_mean": deal_res.usd_irr_mean,
            "usd_irr_median": deal_res.usd_irr_median,
            "usd_irr_p10": deal_res.usd_irr_p10,
            "usd_irr_p90": deal_res.usd_irr_p90,
            "usd_moic_mean": deal_res.usd_moic_mean,
            "p_below_hurdle": deal_res.p_below_hurdle,
            "expected_fx_drag": deal_res.expected_fx_drag,
        }
    ])

    for i, row in scen_data_df.iterrows():
        r_idx = 6 + i
        ws_scen.cell(row=r_idx, column=2, value=row.get("scenario", "")).font = FONT_BOLD
        ws_scen.cell(row=r_idx, column=3, value=row.get("description", "")).font = FONT_REGULAR

        # LCU IRR
        c_lcu = ws_scen.cell(row=r_idx, column=4, value=float(row.get("headline_lcu_irr", 0.0)))
        c_lcu.number_format = "0.0%"

        # Mean USD IRR
        c_mean = ws_scen.cell(row=r_idx, column=5, value=float(row.get("usd_irr_mean", 0.0)))
        c_mean.number_format = "0.0%"
        mean_val = float(row.get("usd_irr_mean", 0.0))
        if mean_val >= deal_config.hurdle_usd_irr:
            c_mean.fill = GREEN_FILL
            c_mean.font = FONT_GREEN
        elif mean_val < 0:
            c_mean.fill = RED_FILL
            c_mean.font = FONT_RED
        else:
            c_mean.fill = AMBER_FILL

        # Median
        c_med = ws_scen.cell(row=r_idx, column=6, value=float(row.get("usd_irr_median", 0.0)))
        c_med.number_format = "0.0%"

        # P10
        c_p10 = ws_scen.cell(row=r_idx, column=7, value=float(row.get("usd_irr_p10", 0.0)))
        c_p10.number_format = "0.0%"

        # P90
        c_p90 = ws_scen.cell(row=r_idx, column=8, value=float(row.get("usd_irr_p90", 0.0)))
        c_p90.number_format = "0.0%"

        # MOIC
        c_moic = ws_scen.cell(row=r_idx, column=9, value=float(row.get("usd_moic_mean", 0.0)))
        c_moic.number_format = '0.00"x"'

        # P below hurdle
        c_pb = ws_scen.cell(row=r_idx, column=10, value=float(row.get("p_below_hurdle", 0.0)))
        c_pb.number_format = "0.0%"
        if float(row.get("p_below_hurdle", 0.0)) > 0.5:
            c_pb.fill = RED_FILL
            c_pb.font = FONT_RED
        else:
            c_pb.fill = GREEN_FILL
            c_pb.font = FONT_GREEN

        # Expected drag
        c_drag = ws_scen.cell(row=r_idx, column=11, value=float(row.get("expected_fx_drag", 0.0)))
        c_drag.number_format = "0.0%"

        for col_idx in range(2, 12):
            ws_scen.cell(row=r_idx, column=col_idx).border = BORDER_ALL

    # Add Bar Chart: Scenario Mean USD IRR vs Target Hurdle
    chart_scen = BarChart()
    chart_scen.type = "col"
    chart_scen.style = 10
    chart_scen.title = f"Expected USD IRR by Scenario (Target Hurdle = {deal_config.hurdle_usd_irr * 100:.1f}%)"
    chart_scen.y_axis.title = "Annualized Return"
    chart_scen.x_axis.title = "Scenario"
    chart_scen.width = 16
    chart_scen.height = 10

    data_ref = Reference(ws_scen, min_col=5, min_row=5, max_row=5 + len(scen_data_df))
    cats_ref = Reference(ws_scen, min_col=2, min_row=6, max_row=5 + len(scen_data_df))
    chart_scen.add_data(data_ref, titles_from_data=True)
    chart_scen.set_categories(cats_ref)
    chart_scen.legend = None
    ws_scen.add_chart(chart_scen, "B16")

    # -------------------------------------------------------------
    # TAB 3: FX SIMULATION & TRAJECTORY
    # -------------------------------------------------------------
    ws_fx = wb.create_sheet(title="FX Simulation Paths")
    ws_fx.views.sheetView[0].showGridLines = True

    ws_fx["B2"] = "Simulated Exchange Rate Trajectory"
    ws_fx["B2"].font = FONT_TITLE
    ws_fx["B3"] = "Monte Carlo simulated path percentiles over investment horizon"
    ws_fx["B3"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    fx_headers = ["Month", "Spot / Mean Rate", "Median (P50)", "P10 (Appreciation/Strong)", "P90 (Severe Devaluation)"]
    for j, h in enumerate(fx_headers, start=2):
        c = ws_fx.cell(row=5, column=j, value=h)
        c.fill = NAVY_HEADER
        c.font = FONT_HEADER_WHITE
        c.alignment = Alignment(horizontal="center")

    max_p = int(deal_df["period"].max())
    n_display_months = max(36, max_p)

    if sim_result is not None:
        paths = sim_result.paths
        n_m = min(n_display_months + 1, paths.shape[1])
        for m in range(n_m):
            row_idx = 6 + m
            col_m = paths[:, m]
            ws_fx.cell(row=row_idx, column=2, value=m).alignment = Alignment(horizontal="center")
            ws_fx.cell(row=row_idx, column=3, value=float(np.mean(col_m))).number_format = "#,##0.00"
            ws_fx.cell(row=row_idx, column=4, value=float(np.median(col_m))).number_format = "#,##0.00"
            ws_fx.cell(row=row_idx, column=5, value=float(np.percentile(col_m, 10))).number_format = "#,##0.00"
            ws_fx.cell(row=row_idx, column=6, value=float(np.percentile(col_m, 90))).number_format = "#,##0.00"
            for col_idx in range(2, 7):
                ws_fx.cell(row=row_idx, column=col_idx).border = BORDER_ALL
    else:
        # Default placeholder trajectory
        for m in range(n_display_months + 1):
            row_idx = 6 + m
            val = 100.0 * ((1 + 0.10) ** (m / 12.0))
            ws_fx.cell(row=row_idx, column=2, value=m)
            ws_fx.cell(row=row_idx, column=3, value=val).number_format = "#,##0.00"
            ws_fx.cell(row=row_idx, column=4, value=val).number_format = "#,##0.00"
            ws_fx.cell(row=row_idx, column=5, value=val * 0.95).number_format = "#,##0.00"
            ws_fx.cell(row=row_idx, column=6, value=val * 1.30).number_format = "#,##0.00"
            for col_idx in range(2, 7):
                ws_fx.cell(row=row_idx, column=col_idx).border = BORDER_ALL

    # Line Chart of FX Paths
    chart_fx = LineChart()
    chart_fx.title = "Simulated Exchange Rate Evolution (P10, Median, P90)"
    chart_fx.style = 13
    chart_fx.y_axis.title = "Exchange Rate (LCU per USD)"
    chart_fx.x_axis.title = "Month"
    chart_fx.width = 17
    chart_fx.height = 10

    fx_data_ref = Reference(ws_fx, min_col=4, min_row=5, max_row=6 + n_display_months, max_col=6)
    fx_cats_ref = Reference(ws_fx, min_col=2, min_row=6, max_row=6 + n_display_months)
    chart_fx.add_data(fx_data_ref, titles_from_data=True)
    chart_fx.set_categories(fx_cats_ref)
    ws_fx.add_chart(chart_fx, "H5")

    # -------------------------------------------------------------
    # TAB 4: CASH FLOW SCHEDULE
    # -------------------------------------------------------------
    ws_cf = wb.create_sheet(title="Cash Flows")
    ws_cf.views.sheetView[0].showGridLines = True

    ws_cf["B2"] = "Deal Cash Flow Schedule"
    ws_cf["B2"].font = FONT_TITLE
    ws_cf["B3"] = f"Repatriation: {deal_config.repatriation_rate} | WHT: {deal_config.withholding_tax_pct * 100:.1f}% | Haircut: {deal_config.blocked_funds_haircut_pct * 100:.1f}%"
    ws_cf["B3"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    cf_headers = ["Period", "Date", "Equity Invested (LCU)", "Distribution (LCU)", "Exit Proceeds (LCU)", "Net Cash Flow (LCU)"]
    for j, h in enumerate(cf_headers, start=2):
        c = ws_cf.cell(row=5, column=j, value=h)
        c.fill = NAVY_HEADER
        c.font = FONT_HEADER_WHITE
        c.alignment = Alignment(horizontal="center")

    for i, row in deal_df.iterrows():
        r_idx = 6 + i
        ws_cf.cell(row=r_idx, column=2, value=int(row["period"])).alignment = Alignment(horizontal="center")
        ws_cf.cell(row=r_idx, column=3, value=str(row["date"]))

        eq = float(row.get("equity_invested_lcu", 0.0))
        dist = float(row.get("distribution_lcu", 0.0))
        exit_p = float(row.get("exit_proceeds_lcu", 0.0))
        net_lcu = dist + exit_p - eq

        c_eq = ws_cf.cell(row=r_idx, column=4, value=eq)
        c_eq.number_format = "#,##0"
        c_dist = ws_cf.cell(row=r_idx, column=5, value=dist)
        c_dist.number_format = "#,##0"
        c_exit = ws_cf.cell(row=r_idx, column=6, value=exit_p)
        c_exit.number_format = "#,##0"

        c_net = ws_cf.cell(row=r_idx, column=7, value=net_lcu)
        c_net.number_format = "#,##0"
        c_net.font = FONT_BOLD
        if net_lcu > 0:
            c_net.font = FONT_GREEN
            c_net.fill = GREEN_FILL
        elif net_lcu < 0:
            c_net.font = FONT_RED

        for col_idx in range(2, 8):
            ws_cf.cell(row=r_idx, column=col_idx).border = BORDER_ALL

    # -------------------------------------------------------------
    # TAB 5: GOVERNANCE & DISCLAIMERS
    # -------------------------------------------------------------
    ws_gov = wb.create_sheet(title="Governance & Disclaimers")
    ws_gov.views.sheetView[0].showGridLines = True

    ws_gov["B2"] = "Model Governance, Methodology & Disclaimers"
    ws_gov["B2"].font = FONT_TITLE

    disclaimers = [
        ("Sparse Historical Episodes", "Frontier market discrete devaluations occur infrequently. Priors dominate model outputs unless long, clean historical panels exist."),
        ("Realized vs Official Rates", "Official pegged or managed rates frequently diverge from parallel market rates or achievable repatriation windows."),
        ("Policy Discretion & Timing", "Devaluations are political decisions by sovereign authorities; the model assesses hazard probability, not exact calendar timing."),
        ("Forward / NDF Illiquidity", "Liquidity in frontier currency forwards and NDFs is often thin. CIP proxy rates illustrate theoretical carry, not firm market depth."),
        ("Not Investment Advice", "Outputs are quantitative stress-testing estimates provided 'as-is' under Apache 2.0 without warranty. Not intended as investment advice."),
    ]

    ws_gov["B4"] = "Topic"
    ws_gov["B4"].font = FONT_HEADER_WHITE
    ws_gov["B4"].fill = NAVY_HEADER
    ws_gov["C4"] = "Governance / Methodological Limitation"
    ws_gov["C4"].font = FONT_HEADER_WHITE
    ws_gov["C4"].fill = NAVY_HEADER

    for i, (top, desc) in enumerate(disclaimers, start=5):
        ws_gov.cell(row=i, column=2, value=top).font = FONT_BOLD
        ws_gov.cell(row=i, column=3, value=desc).font = FONT_REGULAR
        ws_gov.cell(row=i, column=2).border = BORDER_ALL
        ws_gov.cell(row=i, column=3).border = BORDER_ALL

    # Auto-adjust column widths across all worksheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and len(val_str) < 60:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to path
    if output_path is None:
        output_path = Path("reports/deal.xlsx")
    else:
        output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def save_report(
    deal_res: DealAnalysisResult,
    deal_config: DealConfig,
    dev_prob_12m: float,
    dev_prob_24m: float,
    risk_drivers: List[Tuple[str, float, str]],
    deal_df: pd.DataFrame,
    output_path: Union[str, Path],
    hedge_res: Optional[HedgeAnalysisResult] = None,
    comparison_df: Optional[pd.DataFrame] = None,
    sim_result: Optional[SimulationResult] = None,
    as_of_date: Optional[Union[str, dt.date]] = None,
    country: Optional[str] = None,
    format_type: Optional[str] = None,
) -> Path:
    """Save an investment report in either Markdown (.md) or Excel (.xlsx) format.

    Parameters
    ----------
    deal_res : DealAnalysisResult
        Results from deal analysis.
    deal_config : DealConfig
        Deal configuration.
    dev_prob_12m : float
        12-month devaluation probability.
    dev_prob_24m : float
        24-month devaluation probability.
    risk_drivers : List[Tuple[str, float, str]]
        Top risk drivers.
    deal_df : pd.DataFrame
        Deal cash flow schedule.
    output_path : Union[str, Path]
        Target file path.
    hedge_res : Optional[HedgeAnalysisResult], default None
        Hedging comparison.
    comparison_df : Optional[pd.DataFrame], default None
        Scenario comparison.
    sim_result : Optional[SimulationResult], default None
        Simulation paths.
    as_of_date : Optional[Union[str, dt.date]], default None
        As-of date.
    country : Optional[str], default None
        Country name.
    format_type : Optional[str], default None
        Explicit format 'excel' or 'markdown'. If None, detected from file suffix.

    Returns
    -------
    Path
        Path to the saved report.
    """
    out_p = Path(output_path)
    is_excel = (
        (format_type and format_type.lower() in ("excel", "xlsx"))
        or out_p.suffix.lower() == ".xlsx"
    )

    if is_excel:
        return generate_excel_report(
            deal_res=deal_res,
            deal_config=deal_config,
            dev_prob_12m=dev_prob_12m,
            dev_prob_24m=dev_prob_24m,
            risk_drivers=risk_drivers,
            deal_df=deal_df,
            hedge_res=hedge_res,
            comparison_df=comparison_df,
            sim_result=sim_result,
            as_of_date=as_of_date,
            country=country,
            output_path=out_p,
        )
    else:
        md_text = generate_markdown_report(
            deal_res=deal_res,
            deal_config=deal_config,
            dev_prob_12m=dev_prob_12m,
            dev_prob_24m=dev_prob_24m,
            risk_drivers=risk_drivers,
            hedge_res=hedge_res,
            comparison_df=comparison_df,
            as_of_date=as_of_date,
            country=country,
        )
        out_p.parent.mkdir(parents=True, exist_ok=True)
        with open(out_p, "w", encoding="utf-8") as f:
            f.write(md_text)
        return out_p
